# Gera data.js da pagina de planejamentos (rotina-comercial-site/planejamentos), separada da Base Sniper.
# Entram todos os planejamentos da extracao; faixa pelo valor de investimentos fora (S = sem valor fora preenchido).
# Fonte: so 'planejamento financeiro preenchido.xlsx' (extracao do FP da XP, aba result).
# Uso (na pasta Claude): python rotina-comercial-site/planejamentos/gerar_dados.py .
#
# Equipe de cada assessor segue a mesma regra da Base Sniper:
#  1. assessor no ADVISORS do rotina-comercial-site/index.html (casando pelo nome) -> equipe dele;
#  2. lider no LIDER_MAP do dash-1x1-lideres/index.html -> equipe que ele lidera;
#  3. codigo da propria mesa proprietaria -> Mesa proprietaria;
#  4. resto (gestores, codigos sem cadastro) -> FORA_ROTINA.
# Nome do assessor vem do Equipes.xlsx.
import openpyxl, json, sys, unicodedata
from collections import defaultdict

base = sys.argv[1]
site = base + '/rotina-comercial-site'

def norm(s):
    return unicodedata.normalize('NFD', s).encode('ascii', 'ignore').decode().lower().strip()

def js_value(path, var):
    html = open(path, encoding='utf-8').read()
    i = html.index(var + ' = ') + len(var + ' = ')
    return json.JSONDecoder().raw_decode(html[i:].replace('\n', ' '))[0]

def num(v):
    try: return float(v)
    except (TypeError, ValueError): return 0.0

advisors_site = {norm(a['nome']): a['equipe'] for a in js_value(site + '/index.html', 'var ADVISORS')}
html1x1 = open(site + '/dash-1x1-lideres/index.html', encoding='utf-8').read()
i = html1x1.index('LIDER_MAP = ') + len('LIDER_MAP = ')
lideres = json.loads(html1x1[i:html1x1.index('}', i) + 1].replace(',\n  }', '}').replace(',\n}', '}'))

wb = openpyxl.load_workbook(base + '/Equipes.xlsx', data_only=True)
planilha = {r[0]: {'n': r[1], 'e': r[2]} for r in wb['Equipes'].iter_rows(min_row=2, values_only=True) if r[0]}

# ---- Planejamento financeiro: patrimonio declarado no FP (so a planilha, sem Louro) ----
# Assessor = quem criou o FP (CREATED_BY = NOME.Axxxx); se nao tiver codigo A, quem atualizou por ultimo (UPDATED_BY).
ws = openpyxl.load_workbook(base + '/planejamento financeiro preenchido.xlsx', read_only=True, data_only=True)['result']
it = ws.iter_rows(values_only=True); H = list(next(it)); ix = {h: i for i, h in enumerate(H)}
fp_por_conta = {}
for r in it:
    c = str(r[ix['COD_CONTA']] or '').strip()
    if not c.isdigit(): continue
    if c in fp_por_conta and str(fp_por_conta[c][ix['UPDATED_AT']]) >= str(r[ix['UPDATED_AT']]): continue
    fp_por_conta[c] = r
def cod_assessor(r):
    for k in ('CREATED_BY', 'UPDATED_BY'):
        a = str(r[ix[k]] or '').split('.')[-1].upper()
        if a[:1] == 'A' and a[1:].isdigit(): return a
    return 'Sem assessor'
# mesmo cliente com FP em mais de uma conta (a XP repete o patrimonio em cada conta): mesma data de nascimento
# e mesmo assessor -> uma linha so, com os valores do FP mais recente
por_cliente = defaultdict(list)
for c, r in fp_por_conta.items():
    nasc = str(r[ix['DATA_NASCIMENTO']] or '')
    k = (nasc, cod_assessor(r)) if nasc not in ('', 'null', 'None') else c
    por_cliente[k].append(c)
fp, ref_fp_xls = [], max(str(r[ix['DAT_EXPORTACAO']])[:10] for r in fp_por_conta.values())
for cs in por_cliente.values():
    cs.sort(key=lambda k: str(fp_por_conta[k][ix['UPDATED_AT']]), reverse=True)
    c, r = cs[0], fp_por_conta[cs[0]]
    e = num(r[ix['VAL_INV_EXTERNO']]) + num(r[ix['VAL_PREV']])   # patrimonio fora da XP = investimentos fora + previdencia
    fp.append({'c': c, 'a': cod_assessor(r), 's': r[ix['SEGMENTO_COMERCIAL']] or '', 'f': 'A' if e >= 1e6 else 'B' if e >= 3e5 else 'C' if e >= 1e5 else 'D' if e >= 1 else 'S',
               'e': round(e), 'im': round(num(r[ix['VAL_IMOVEIS']])), 'vc': round(num(r[ix['VAL_VEICULOS']])),
               'am': round(num(r[ix['VAL_APORTE_MENSAL']])), 'op': 1 if num(r[ix['QTD_OPIN']]) > 0 else 0,
               'u': str(r[ix['UPDATED_AT']])[:10], 'sc': str(r[ix['FAIXA_SCORE_PREENCHIMENTO']] or '').split('. ')[-1].replace('null', ''), 'o': cs[1:]})

FORA_ROTINA = 'Fora da Rotina Comercial'
equipe_do_lider = {}
for equipe, lider in lideres.items():
    equipe_do_lider.setdefault(norm(lider), equipe)  # George: Sparta (1a do mapa) antes da Mesa
equipes_rotina = set(advisors_site.values())

assessores = {}
for a in {r['a'] for r in fp}:
    p = planilha.get(a, {'n': a, 'e': ''})
    nome = p['n'] or a
    e = advisors_site.get(norm(nome)) or equipe_do_lider.get(norm(nome))
    if not e and p['e'] == 'Mesa proprietária' and p['e'] in equipes_rotina:
        e = p['e']
    assessores[a] = {'n': nome, 'e': e or FORA_ROTINA}

data = {'refFpXls': ref_fp_xls, 'assessores': assessores, 'lideres': lideres, 'fp': fp}
open(site + '/planejamentos/data.js', 'w', encoding='utf-8').write('window.FP_DATA = ' + json.dumps(data, ensure_ascii=False) + ';\n')
print(len(fp), 'clientes,', len(assessores), 'assessores')
