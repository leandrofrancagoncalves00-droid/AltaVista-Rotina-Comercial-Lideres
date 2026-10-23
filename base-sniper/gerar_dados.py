# Gera data.js a partir de Base_Sniper_Alta_Vista.xlsx, Alta_Vista_Open_Investments_Analise.xlsx
# e do CSV mais recente Avenue_contas_AAAA-MM-DD.csv (extraido do harness Louro, posicao_diaria da Avenue)
# e de 'clientes afastados.xlsx' (forca-tarefa; as abas por time da planilha sao ignoradas: a equipe
# vem do codigo do assessor, pela mesma regra das outras bases).
# Uso (na pasta Claude): python rotina-comercial-site/base-sniper/gerar_dados.py .
#
# Equipe de cada assessor segue a Rotina Comercial:
#  1. assessor no ADVISORS do rotina-comercial-site/index.html (casando pelo nome) -> equipe dele;
#  2. lider no LIDER_MAP do dash-1x1-lideres/index.html -> equipe que ele lidera;
#  3. codigo da propria mesa proprietaria -> Mesa proprietaria;
#  4. resto (gestores, codigos sem cadastro) -> FORA_ROTINA.
# Nome do assessor vem do Equipes.xlsx.
import openpyxl, json, sys, unicodedata, csv, glob, os
from collections import defaultdict

base = sys.argv[1]
site = base + '/rotina-comercial-site'

def norm(s):
    return unicodedata.normalize('NFD', s).encode('ascii', 'ignore').decode().lower().strip()

def js_value(path, var):
    html = open(path, encoding='utf-8').read()
    i = html.index(var + ' = ') + len(var + ' = ')
    return json.JSONDecoder().raw_decode(html[i:].replace('\n', ' '))[0]

advisors_site = {norm(a['nome']): a['equipe'] for a in js_value(site + '/index.html', 'var ADVISORS')}
html1x1 = open(site + '/dash-1x1-lideres/index.html', encoding='utf-8').read()
i = html1x1.index('LIDER_MAP = ') + len('LIDER_MAP = ')
lideres = json.loads(html1x1[i:html1x1.index('}', i) + 1].replace(',\n  }', '}').replace(',\n}', '}'))

wb = openpyxl.load_workbook(base + '/Equipes.xlsx', data_only=True)
planilha = {r[0]: {'n': r[1], 'e': r[2]} for r in wb['Equipes'].iter_rows(min_row=2, values_only=True) if r[0]}

# ---- Base Sniper (P1 + P2) ----
s = openpyxl.load_workbook(base + '/Base_Sniper_Alta_Vista.xlsx', data_only=True)
sniper = []
for t in ['Alta Vista P1', 'Alta Vista P2']:
    for r in s[t].iter_rows(min_row=2, values_only=True):
        if not r[2]: continue
        sniper.append({'p': r[0], 'f': r[1], 'c': str(r[2]), 'a': r[3], 's': r[4],
                       'x': round(r[5] or 0), 'l': round(r[6] or 0), 'w': round(r[7] or 0, 4), 'g': round(r[8] or 0)})

# ---- Open Investments ----
o = openpyxl.load_workbook(base + '/Alta_Vista_Open_Investments_Analise.xlsx', data_only=True)
inst = defaultdict(lambda: defaultdict(float))
for r in o['Base tratada'].iter_rows(min_row=3, values_only=True):
    if r[1] is None or not isinstance(r[9], (int, float)): continue
    inst[str(r[1])][r[2]] += r[9] or 0
opin = []
for r in o['Contas'].iter_rows(min_row=4, values_only=True):
    if not isinstance(r[1], (int, float)): continue
    c = str(r[1])
    top = sorted(inst[c].items(), key=lambda kv: -kv[1])[:5]
    opin.append({'c': c, 'a': r[2], 'f': r[3][0], 'fl': r[3], 't': round(r[4] or 0), 'ni': r[5], 'ip': r[6], 'pp': round(r[7] or 0, 3),
                 'gk': r[8] or '', 'ab': r[9] or '',
                 'mix': [round(v or 0) for v in r[10:15]], 'v90': round(r[15] or 0), 'v12': round(r[16] or 0),
                 'ac': round(r[17] or 0), 'inst': [[k, round(v)] for k, v in top]})

# ---- Avenue ----
av_file = sorted(glob.glob(base + '/Avenue_contas_*.csv'))[-1]
ref_avenue = os.path.basename(av_file)[len('Avenue_contas_'):-4]
avenue, nomes_db = [], {}
for r in csv.DictReader(open(av_file, encoding='utf-8')):
    t = float(r['net'])
    if t < 1: continue
    f = 'A' if t >= 1e6 else 'B' if t >= 3e5 else 'C' if t >= 1e5 else 'D'
    avenue.append({'c': r['conta'], 'a': r['cod_assessor'], 'f': f, 't': round(t),
                   'mix': [round(float(r[k])) for k in ('fu', 'alt', 'rv', 'rf', 'caixa')], 'ab': r['abertura']})
    nomes_db[r['cod_assessor']] = r['assessor']

# ---- Clientes afastados ----
fa = openpyxl.load_workbook(base + '/clientes afastados.xlsx', data_only=True)
afastados, vistos = [], set()
for ws in fa.worksheets:
    if ws.title == 'Resumo': continue
    for r in ws.iter_rows(min_row=5, values_only=True):
        if len(r) < 13 or not r[10] or not str(r[10]).strip().isdigit() or str(r[10]) in vistos: continue
        vistos.add(str(r[10]))
        afastados.append({'c': str(r[10]).strip(), 'k': str(r[11]).strip(), 'a': str(r[12]).strip()})

FORA_ROTINA = 'Fora da Rotina Comercial'
equipe_do_lider = {}
for equipe, lider in lideres.items():
    equipe_do_lider.setdefault(norm(lider), equipe)  # George: Sparta (1a do mapa) antes da Mesa
equipes_rotina = set(advisors_site.values())

assessores = {}
for a in {r['a'] for r in sniper} | {r['a'] for r in opin} | {r['a'] for r in avenue} | {r['a'] for r in afastados}:
    p = planilha.get(a, {'n': nomes_db.get(a, a), 'e': ''})
    nome = p['n'] or a
    e = advisors_site.get(norm(nome)) or equipe_do_lider.get(norm(nome))
    if not e and p['e'] == 'Mesa proprietária' and p['e'] in equipes_rotina:
        e = p['e']
    assessores[a] = {'n': nome, 'e': e or FORA_ROTINA}

data = {'refSniper': 'set/2026', 'refAvenue': ref_avenue, 'assessores': assessores, 'lideres': lideres, 'sniper': sniper, 'opin': opin, 'avenue': avenue, 'afastados': afastados}
open(site + '/base-sniper/data.js', 'w', encoding='utf-8').write('window.BASES_DATA = ' + json.dumps(data, ensure_ascii=False) + ';\n')
print(len(sniper), 'sniper,', len(opin), 'opin,', len(avenue), 'avenue,', len(afastados), 'afastados,', len(assessores), 'assessores')
